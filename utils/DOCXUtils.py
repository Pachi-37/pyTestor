import os

import pandas
from docx.table import _Cell, _Row
from docx.shared import Inches

from utils.ImageUtils import ImageUtils


class XlsxData:
    def __init__(self):
        pass


class DOCXUtils:
    @staticmethod
    def replace_string_in_docx_two(cell: _Cell, community, number):
        current_text = cell.text.format(community, number)
        cell.paragraphs[0].clear()
        run = cell.paragraphs[0].add_run()
        run.text = current_text
        return current_text + ".docx"

    @staticmethod
    def replace_string_in_docx(cell: _Cell, info):
        if pandas.isna(info):
            info = "/"
        current_text = cell.text.format(info)
        cell.paragraphs[0].clear()
        run = cell.paragraphs[0].add_run()
        run.text = current_text
        return current_text + ".docx"

    @staticmethod
    def insert_image_in_docx(row: _Row, cell: _Cell, image_path, percentage=100):
        resize_image = ImageUtils.resize_image(image_path, percentage)
        width, height = ImageUtils.get_image_size_inch(resize_image)
        cell.width = Inches(width)
        row.height = Inches(height)
        cell.paragraphs[0].add_run().add_picture(
            os.path.join(os.path.dirname(image_path), resize_image))
